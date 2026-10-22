#!/usr/bin/env python3
"""Render a JSON document spec into a DOCX that follows the UniCredit document template.

The spec format is in references/spec-format.md. Every visual value comes from
assets/doc-tokens.json, so a design change is a token change, not a code change.

Exit codes: 0 written, 1 invalid spec, 2 usage or I/O error.
"""
import argparse
import json
import re
import sys
import tempfile
import zipfile
from pathlib import Path

try:
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Emu, Pt, RGBColor, Twips
except ImportError:
    print("python-docx is required: pip install python-docx", file=sys.stderr)
    sys.exit(2)

HERE = Path(__file__).resolve().parent
TOKENS = json.loads((HERE.parent / "assets" / "doc-tokens.json").read_text(encoding="utf-8"))
FONT = TOKENS["font"]["family"]
COLORS = TOKENS["colors"]
INLINE = re.compile(r"(\*\*.+?\*\*|\[TBC\]|\[Proposed\])")
ROMAN = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X"]
FRONT_STYLE = "UC Front Subheading"
# Set by add_bullet_numbering. Ids are the next free ones, not a fixed high value: simple
# renderers (macOS Quick Look, mail previews) resolve abstractNumId by position and drop the bullets otherwise.
BULLET_NUM_ID = None


class SpecError(Exception):
    pass


def color(name):
    return RGBColor.from_string(COLORS.get(name, name))


# ---------------------------------------------------------------- low-level XML helpers

def force_fonts(root):
    """Pin every rFonts under root to the template font and drop theme fonts, which would win."""
    for rf in root.iter(qn("w:rFonts")):
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            rf.attrib.pop(qn(attr), None)
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rf.set(qn(attr), FONT)


def set_spacing(ppr, before=None, after=None, line=None):
    sp = ppr.find(qn("w:spacing"))
    if sp is None:
        sp = OxmlElement("w:spacing")
        ppr.append(sp)
    if before is not None:
        sp.set(qn("w:before"), str(before))
    if after is not None:
        sp.set(qn("w:after"), str(after))
    if line is not None:
        sp.set(qn("w:line"), str(line))
        sp.set(qn("w:lineRule"), "auto")


def para_format(p, before=None, after=None, line=None, keep_next=False, page_break_before=False):
    pf = p.paragraph_format
    if keep_next:
        pf.keep_with_next = True
    if page_break_before:
        pf.page_break_before = True
    set_spacing(p._p.get_or_add_pPr(), before, after, line)


def style_run(run, size=None, bold=None, italic=None, color_name="text"):
    run.font.name = FONT
    rpr = run._r.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    rf.set(qn("w:eastAsia"), FONT)
    rf.set(qn("w:cs"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color_name:
        run.font.color.rgb = color(color_name)


def add_inline(p, text, size=None, bold=False, italic=False, color_name="text"):
    """Write text into p, honouring **bold** and the [TBC] / [Proposed] tags."""
    tag = TOKENS["tag"]
    for chunk in INLINE.split(text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**"):
            style_run(p.add_run(chunk[2:-2]), size, True, italic, color_name)
        elif chunk in tag["values"]:
            style_run(p.add_run(chunk), size, tag["bold"], italic, tag["color"])
        else:
            style_run(p.add_run(chunk), size, bold, italic, color_name)


# ---------------------------------------------------------------- document setup

def setup_document(doc, spec):
    page = TOKENS["page"]
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Twips(page["width"]), Twips(page["height"])
    sec.top_margin, sec.bottom_margin = Twips(page["margin_top"]), Twips(page["margin_bottom"])
    sec.left_margin, sec.right_margin = Twips(page["margin_left"]), Twips(page["margin_right"])
    sec.header_distance, sec.footer_distance = Twips(page["header_distance"]), Twips(page["footer_distance"])

    styles_el = doc.styles.element
    force_fonts(styles_el)
    rpr_default = styles_el.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
    lang = rpr_default.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr_default.append(lang)
    lang.set(qn("w:val"), spec.get("language", TOKENS["font"]["language"]))

    body = TOKENS["body"]
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(body["size"])
    normal.font.color.rgb = color("text")
    set_spacing(normal.element.get_or_add_pPr(), body["space_before"], body["space_after"], body["line"])

    for name, tok in (("Heading 1", TOKENS["headings"]["h1"]), ("Heading 2", TOKENS["headings"]["h2"])):
        st = doc.styles[name]
        st.font.name, st.font.size, st.font.bold, st.font.italic = FONT, Pt(tok["size"]), tok["bold"], False
        st.font.color.rgb = color("text")
        ppr = st.element.get_or_add_pPr()
        set_spacing(ppr, tok["space_before"], tok["space_after"], tok["line"])
        st.paragraph_format.left_indent = Twips(tok["indent_left"])
        st.paragraph_format.first_line_indent = Twips(0)
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.keep_together = True

    # Front-matter subheadings look like Heading 2 but carry no outline level, so the TOC skips them.
    front = doc.styles.add_style(FRONT_STYLE, WD_STYLE_TYPE.PARAGRAPH)
    front.base_style = normal
    h2 = TOKENS["headings"]["h2"]
    front.font.name, front.font.size, front.font.bold = FONT, Pt(h2["size"]), True
    front.font.color.rgb = color("text")
    front.paragraph_format.keep_with_next = True
    set_spacing(front.element.get_or_add_pPr(), h2["space_before"], h2["space_after"], h2["line"])

    toc = TOKENS["toc"]
    for name, tok in (("TOC 1", toc["toc1"]), ("TOC 2", toc["toc2"])):
        try:
            st = doc.styles[name]
        except KeyError:
            st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            st.base_style = normal
        st.font.name, st.font.size, st.font.bold = FONT, Pt(tok["size"]), tok["bold"]
        st.font.color.rgb = color("text")
        st.paragraph_format.left_indent = Twips(tok["indent_left"])
        set_spacing(st.element.get_or_add_pPr(), tok["space_before"], tok["space_after"], TOKENS["body"]["line"])
        st.paragraph_format.tab_stops.add_tab_stop(Twips(page["content_width"]), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)

    force_fonts(styles_el)
    add_bullet_numbering(doc)
    if TOKENS["toc"]["update_fields_on_open"]:
        add_update_fields(doc)

    meta = spec["cover"].get("metadata", {})
    doc.core_properties.title = spec["cover"]["title"]
    doc.core_properties.author = meta.get("Author", "")
    doc.core_properties.comments = ""


def add_bullet_numbering(doc):
    global BULLET_NUM_ID
    b = TOKENS["bullets"]
    numbering = doc.part.numbering_part.element
    abstract_id = 1 + max((int(a.get(qn("w:abstractNumId"))) for a in numbering.findall(qn("w:abstractNum"))), default=-1)
    BULLET_NUM_ID = 1 + max((int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))), default=0)
    absn = OxmlElement("w:abstractNum")
    absn.set(qn("w:abstractNumId"), str(abstract_id))
    mlt = OxmlElement("w:multiLevelType")
    mlt.set(qn("w:val"), "hybridMultilevel")
    absn.append(mlt)
    for ilvl, glyph in ((0, b["glyph_level0"]), (1, b["glyph_level1"])):
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), str(ilvl))
        for tag, val in (("w:start", "1"), ("w:numFmt", "bullet"), ("w:lvlText", glyph), ("w:lvlJc", "left")):
            e = OxmlElement(tag)
            e.set(qn("w:val"), val)
            lvl.append(e)
        ppr = OxmlElement("w:pPr")
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), str(b["indent_left"] * (ilvl + 1)))
        ind.set(qn("w:hanging"), str(b["hanging"]))
        ppr.append(ind)
        lvl.append(ppr)
        rpr = OxmlElement("w:rPr")
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
        lvl.append(rpr)
        absn.append(lvl)
    force_fonts(absn)
    first_num = numbering.find(qn("w:num"))
    if first_num is not None:
        first_num.addprevious(absn)
    else:
        numbering.append(absn)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(BULLET_NUM_ID))
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), str(abstract_id))
    num.append(ref)
    numbering.append(num)


def add_update_fields(doc):
    settings = doc.settings.element
    uf = OxmlElement("w:updateFields")
    uf.set(qn("w:val"), "true")
    # CT_Settings is a sequence: updateFields must precede these siblings or Word reports corruption.
    later = ("w:hdrShapeDefaults", "w:footnotePr", "w:endnotePr", "w:compat", "w:docVars", "w:rsids",
             "m:mathPr", "w:attachedSchema", "w:themeFontLang", "w:clrSchemeMapping", "w:doNotIncludeSubdocsInStats",
             "w:doNotAutoCompressPictures", "w:forceUpgrade", "w:captions", "w:readModeInkLockDown",
             "w:smartTagType", "w:shapeDefaults", "w:doNotEmbedSmartTags", "w:decimalSymbol", "w:listSeparator")
    later_q = {qn(t) for t in later}
    for child in settings:
        if child.tag in later_q:
            child.addprevious(uf)
            return
    settings.append(uf)


# ---------------------------------------------------------------- blocks

def add_paragraph(doc, text, **fmt):
    p = doc.add_paragraph()
    body = TOKENS["body"]
    para_format(p, fmt.get("before", body["space_before"]), fmt.get("after", body["space_after"]), body["line"])
    add_inline(p, text, size=fmt.get("size"), italic=fmt.get("italic", False), color_name=fmt.get("color", "text"))
    return p


def add_bullets(doc, items, level=0):
    b = TOKENS["bullets"]
    for item in items:
        if isinstance(item, list):
            add_bullets(doc, item, level + 1)
            continue
        p = doc.add_paragraph(style="List Paragraph")
        ppr = p._p.get_or_add_pPr()
        numpr = OxmlElement("w:numPr")
        il = OxmlElement("w:ilvl")
        il.set(qn("w:val"), str(level))
        ni = OxmlElement("w:numId")
        ni.set(qn("w:val"), str(BULLET_NUM_ID))
        numpr.append(il)
        numpr.append(ni)
        ppr.insert(1 if ppr.find(qn("w:pStyle")) is not None else 0, numpr)
        set_spacing(ppr, b["space_before"], b["space_after"], b["line"])
        add_inline(p, item)


def add_note(doc, text):
    n = TOKENS["note"]
    p = doc.add_paragraph()
    para_format(p, n["space_before"], n["space_after"], TOKENS["body"]["line"])
    style_run(p.add_run(n["lead"] + " "), bold=True)
    add_inline(p, text)


def add_legend(doc, text):
    lg = TOKENS["legend"]
    p = doc.add_paragraph()
    para_format(p, lg["space_before"], lg["space_after"], TOKENS["body"]["line"])
    add_inline(p, text, size=lg["size"], color_name=lg["color"])


def column_widths(columns, rows, widths_cm, size):
    total = TOKENS["table"]["width"]
    if widths_cm:
        if len(widths_cm) != len(columns):
            raise SpecError(f"table {columns}: widths_cm has {len(widths_cm)} values for {len(columns)} columns")
        raw = [w * 567 for w in widths_cm]
        scale = total / sum(raw)
        widths = [int(w * scale) for w in raw]
    else:
        # Every column first gets room for its longest word, so an ID such as "N1" never wraps;
        # the rest is shared by average content length, capped so one long cell cannot starve the others.
        char_tw = 0.6 * size * 20  # an Arial character averages about 0.55 em; 0.6 leaves room for bold
        pad = TOKENS["table"]["cell_margin_left"] + TOKENS["table"]["cell_margin_right"] + 60
        mins, weights = [], []
        for i, col in enumerate(columns):
            cells = [_plain(r[i]) for r in rows]
            longest = max(len(w) for w in " ".join([col] + cells).split()) if (col or any(cells)) else 1
            mins.append(max(int(longest * char_tw + pad), 567))
            weights.append(max(min(sum(len(c) for c in cells) / max(len(cells), 1), 60), len(col), 1))
        spare = total - sum(mins)
        if spare < 0:
            raise SpecError(f"table {columns}: too many columns for the page; split it or pass widths_cm")
        widths = [int(m + spare * w / sum(weights)) for m, w in zip(mins, weights)]
    widths[-1] += total - sum(widths)
    return widths


def _plain(value):
    text = " ".join(value) if isinstance(value, list) else str(value)
    return text.replace("**", "")


def cell_borders(tcpr):
    t = TOKENS["table"]
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(t["border_eighths"]))
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), COLORS["table_border"])
        borders.append(e)
    tcpr.append(borders)


def add_table(doc, block):
    t = TOKENS["table"]
    columns = block["columns"]
    rows = block.get("rows", [])
    size_key = block.get("size", "compact")
    if size_key not in ("standard", "compact"):
        raise SpecError(f"table {columns}: size must be standard or compact")
    size, line = t[size_key]["size"], t[size_key]["line"]
    first_bold = block.get("first_column_bold", False)
    for r in rows:
        if len(r) != len(columns):
            raise SpecError(f"table {columns}: row {r} has {len(r)} cells")
    widths = column_widths(columns, rows, block.get("widths_cm"), size)

    table = doc.add_table(rows=1 + len(rows), cols=len(columns))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tblpr = table._tbl.tblPr
    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.append(tblw)
    tblw.set(qn("w:w"), str(t["width"]))
    tblw.set(qn("w:type"), "dxa")
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblpr.append(layout)
    for gc, w in zip(table._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(w))

    for ri, row in enumerate(table.rows):
        trpr = row._tr.get_or_add_trPr()
        if t["cant_split_rows"]:
            trpr.append(OxmlElement("w:cantSplit"))
        if ri == 0 and t["repeat_header"]:
            trpr.append(OxmlElement("w:tblHeader"))
        values = columns if ri == 0 else rows[ri - 1]
        for ci, (cell, value) in enumerate(zip(row.cells, values)):
            cell.width = Twips(widths[ci])
            tcpr = cell._tc.get_or_add_tcPr()
            cell_borders(tcpr)
            if ri == 0:
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:color"), "auto")
                shd.set(qn("w:fill"), COLORS["table_header_fill"])
                tcpr.append(shd)
            # Per cell, as in the reference: a table-level tblCellMar is ignored by some renderers.
            mar = OxmlElement("w:tcMar")
            for side in ("top", "left", "bottom", "right"):
                e = OxmlElement(f"w:{side}")
                e.set(qn("w:w"), str(t[f"cell_margin_{side}"]))
                e.set(qn("w:type"), "dxa")
                mar.append(e)
            tcpr.append(mar)
            parts = value if isinstance(value, list) else [str(value)]
            for pi, text in enumerate(parts):
                p = cell.paragraphs[0] if pi == 0 else cell.add_paragraph()
                para_format(p, t["cell_space_before"], t["cell_space_after"], line)
                add_inline(p, text, size=size, bold=(ri == 0 or (first_bold and ci == 0)))
    return table


def add_blocks(doc, blocks):
    for blk in blocks:
        kind = blk.get("type")
        if kind == "paragraph":
            add_paragraph(doc, blk["text"])
        elif kind == "bullets":
            add_bullets(doc, blk["items"])
        elif kind == "table":
            add_table(doc, blk)
        elif kind == "note":
            add_note(doc, blk["text"])
        elif kind == "legend":
            add_legend(doc, blk["text"])
        else:
            raise SpecError(f"unknown block type {kind!r}")


# ---------------------------------------------------------------- page furniture

def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def front_title(doc, text):
    ft = TOKENS["front_title"]
    p = doc.add_paragraph()
    para_format(p, 0, ft["space_after"], TOKENS["body"]["line"], keep_next=True)
    style_run(p.add_run(text), ft["size"], ft["bold"])


def front_subheading(doc, text):
    p = doc.add_paragraph(style=FRONT_STYLE)
    style_run(p.add_run(text), bold=True)


def part_label(doc, text):
    pl = TOKENS["part_label"]
    p = doc.add_paragraph()
    para_format(p, pl["space_before"], pl["space_after"], TOKENS["body"]["line"], keep_next=True,
                page_break_before=pl["page_break_before"])
    ppr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(pl["border_bottom_eighths"]))
    bottom.set(qn("w:space"), str(pl["border_space_pt"]))
    bottom.set(qn("w:color"), COLORS[pl["color"]])
    bdr.append(bottom)
    # pBdr sits after keepNext/pageBreakBefore and before spacing in the CT_PPr sequence.
    ppr.find(qn("w:spacing")).addprevious(bdr)
    style_run(p.add_run(text), pl["size"], pl["bold"], color_name=pl["color"])


def heading(doc, text, level):
    p = doc.add_paragraph(style=f"Heading {level}")
    style_run(p.add_run(text), bold=True)
    return p


def resolve_logo(args):
    if args.logo:
        path = Path(args.logo)
        if not path.is_file():
            raise FileNotFoundError(f"logo not found: {path}")
        return path
    if args.logo_from:
        with zipfile.ZipFile(args.logo_from) as z:
            media = sorted(n for n in z.namelist() if n.startswith("word/media/") and n.lower().endswith((".png", ".jpg", ".jpeg")))
            if not media:
                raise FileNotFoundError(f"no image in {args.logo_from}")
            out = Path(tempfile.mkdtemp()) / Path(media[0]).name
            out.write_bytes(z.read(media[0]))
            print(f"logo taken from {args.logo_from}:{media[0]}", file=sys.stderr)
            return out
    return None


def cover(doc, spec, logo):
    c = TOKENS["cover"]
    cv = spec["cover"]
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para_format(p, 0, c["logo_space_after"])
    if logo:
        p.add_run().add_picture(str(logo), width=Emu(c["logo_width_emu"]))
    else:
        style_run(p.add_run("[UniCredit logo]"), 14, True, color_name="secondary")
        print("warning: no --logo / --logo-from given, cover carries a placeholder", file=sys.stderr)
    for _ in range(c["blank_lines_after_logo"]):
        para_format(doc.add_paragraph(), 0, 0, TOKENS["body"]["line"])

    for key in ("title", "subtitle", "abstract"):
        if key not in cv:
            if key == "title":
                raise SpecError("cover.title is required")
            continue
        tok = c[key]
        p = doc.add_paragraph()
        para_format(p, 0, tok["space_after"], TOKENS["body"]["line"], keep_next=(key == "title"))
        add_inline(p, cv[key], size=tok["size"], bold=tok.get("bold", False), italic=tok.get("italic", False),
                   color_name=tok.get("color", "text"))

    meta = cv.get("metadata", {})
    cols = c["metadata_columns"]
    missing = [k for k in cols if k not in meta]
    if missing:
        raise SpecError(f"cover.metadata misses {missing}")
    add_table(doc, {"columns": cols, "rows": [[meta[k] for k in cols]], "size": "standard",
                    "widths_cm": c["metadata_widths_cm"]})


def document_control(doc, spec):
    dc = spec["document_control"]
    page_break(doc)
    front_title(doc, "Document control")
    front_subheading(doc, "Revision history")
    add_table(doc, {"columns": ["Version", "Date", "Description"], "rows": dc["revisions"], "size": "standard",
                    "widths_cm": [2.19, 2.92, 12.48]})
    front_subheading(doc, "Approval")
    add_table(doc, {"columns": ["Role", "Organisation", "Approval scope"], "rows": dc["approvals"], "size": "standard",
                    "widths_cm": [5.11, 3.28, 9.2]})
    front_subheading(doc, "Conventions")
    add_bullets(doc, dc["conventions"])


def toc(doc, entries):
    tk = TOKENS["toc"]
    page_break(doc)
    front_title(doc, tk["title"])
    # A complex field spanning paragraphs: begin and separate in the first, end in the last.
    # The entries between are a pre-rendered result, so previews show the outline before Word updates it.
    def fld(run, kind):
        e = OxmlElement("w:fldChar")
        e.set(qn("w:fldCharType"), kind)
        run._r.append(e)

    if not entries:
        entries = [(1, "Right-click and choose Update Field to build the table of contents.")]
    for i, (level, text) in enumerate(entries):
        p = doc.add_paragraph(style=f"TOC {level}")
        if i == 0:
            fld(p.add_run(), "begin")
            r = p.add_run()
            instr = OxmlElement("w:instrText")
            instr.set(qn("xml:space"), "preserve")
            instr.text = f" {tk['instruction']} "
            r._r.append(instr)
            fld(p.add_run(), "separate")
        style_run(p.add_run(text), color_name="text")
        if i == len(entries) - 1:
            fld(p.add_run(), "end")


# ---------------------------------------------------------------- driver

def outline(spec):
    entries, n = [], 0
    for part in spec.get("parts", []):
        for sec in part.get("sections", []):
            n += 1
            entries.append((1, f"{n}. {sec['title']}"))
            for m, sub in enumerate(sec.get("subsections", []), 1):
                entries.append((2, f"{n}.{m} {sub['title']}"))
    for i, ann in enumerate(spec.get("annexes", [])):
        letter = chr(ord("A") + i)
        entries.append((1, f"Annex {letter}. {ann['title']}"))
        for m, sub in enumerate(ann.get("subsections", []), 1):
            entries.append((2, f"{letter}.{m} {sub['title']}"))
    return entries


def build(spec, out, logo):
    for key in ("cover", "document_control", "parts"):
        if key not in spec:
            raise SpecError(f"missing top-level key {key!r}")
    doc = Document()
    setup_document(doc, spec)
    cover(doc, spec, logo)
    document_control(doc, spec)
    toc(doc, outline(spec))

    n = 0
    for pi, part in enumerate(spec["parts"]):
        part_label(doc, f"PART {ROMAN[pi]}: {part['label'].upper()}")
        for sec in part.get("sections", []):
            n += 1
            heading(doc, f"{n}. {sec['title']}", 1)
            add_blocks(doc, sec.get("blocks", []))
            for m, sub in enumerate(sec.get("subsections", []), 1):
                heading(doc, f"{n}.{m} {sub['title']}", 2)
                add_blocks(doc, sub.get("blocks", []))
    if spec.get("annexes"):
        part_label(doc, "ANNEXES")
        for i, ann in enumerate(spec["annexes"]):
            letter = chr(ord("A") + i)
            heading(doc, f"Annex {letter}. {ann['title']}", 1)
            add_blocks(doc, ann.get("blocks", []))
            for m, sub in enumerate(ann.get("subsections", []), 1):
                heading(doc, f"{letter}.{m} {sub['title']}", 2)
                add_blocks(doc, sub.get("blocks", []))

    doc.save(out)


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("spec", help="JSON spec, see references/spec-format.md")
    ap.add_argument("out", help="output .docx path")
    g = ap.add_mutually_exclusive_group()
    g.add_argument("--logo", help="UniCredit logo image (PNG or JPG)")
    g.add_argument("--logo-from", help="existing UniCredit .docx to take the logo from (first image in word/media)")
    args = ap.parse_args()
    try:
        spec = json.loads(Path(args.spec).read_text(encoding="utf-8"))
        logo = resolve_logo(args)
    except (OSError, json.JSONDecodeError, zipfile.BadZipFile) as e:
        print(f"error: {e}", file=sys.stderr)
        return 2
    try:
        build(spec, args.out, logo)
    except SpecError as e:
        print(f"invalid spec: {e}", file=sys.stderr)
        return 1
    except OSError as e:
        print(f"error: {e}", file=sys.stderr)
        return 2
    print(f"written {args.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
